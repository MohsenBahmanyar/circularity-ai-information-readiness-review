from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"data_excel"/"07_Figure_and_Table_Source_Data.xlsx"
SVG=ROOT/"figures"/"Figure_7_Core_Attribute_AI_Functions.svg"
PNG=ROOT/"figures"/"Figure_7_Core_Attribute_AI_Functions.png"
FUNCTIONS=["Extract/Retrieve","Infer/Complete","Predict/Assess","Diagnose","Recommend","Generate plus Rank/Optimize"]
COLORS={"A":"#24747a","B":"#d38b1f","C":"#d8e8ee","D":"#f1f3f4"}

def f(size,bold=False):
    for name in (["arialbd.ttf","DejaVuSans-Bold.ttf"] if bold else ["arial.ttf","DejaVuSans.ttf"]):
        try:return ImageFont.truetype(name,size)
        except OSError:pass
    return ImageFont.load_default()

def main():
    ws=load_workbook(SOURCE,read_only=True,data_only=True)["Figure_7"]; rows=list(ws.iter_rows(values_only=True)); h=[str(v) for v in rows[2]]; data=[dict(zip(h,r)) for r in rows[3:] if r[0]]
    attrs=[]; by={}; bands={}
    for r in data:
        if r["attribute"] not in by: attrs.append(r["attribute"]); by[r["attribute"]]={}; bands[r["attribute"]]=r["readiness_band"]
        by[r["attribute"]][r["ai_function"]]=r
    assert len(attrs)==19 and Counter(bands.values())==Counter({"Band I":8,"Band II":7,"Band III":4})
    band_order={"Band I":0,"Band II":1,"Band III":2}
    attrs.sort(key=lambda name:(band_order[bands[name]],str(name)))
    width,height=3500,2650; image=Image.new("RGB",(width,height),"white");d=ImageDraw.Draw(image);svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#fff"/>']
    title="AI functions demonstrated for source-grounded circularity attributes";d.text((120,45),title,font=f(54,True),fill="#173b4f");svg.append(f'<text x="120" y="105" font-family="Arial" font-size="54" font-weight="700" fill="#173b4f">{escape(title)}</text>')
    subtitle="Eight source-grounded attributes intersect with verified early-design AI context";d.text((120,120),subtitle,font=f(25,True),fill="#24747a");svg.append(f'<text x="120" y="155" font-family="Arial" font-size="25" font-weight="700" fill="#24747a">{escape(subtitle)}</text>')
    left,top,attr_w,cell_w,row_h,band_h=120,220,1060,365,82,65;y=top
    d.rectangle((left,y,left+attr_w,y+90),fill="#173b4f");d.text((left+18,y+28),"Source-grounded attribute",font=f(22,True),fill="white");svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="90" fill="#173b4f"/><text x="{left+18}" y="{y+56}" font-family="Arial" font-size="22" font-weight="700" fill="#fff">Source-grounded attribute</text>')
    for i,fn in enumerate(FUNCTIONS):
        x=left+attr_w+i*cell_w;label=fn.replace("Generate plus Rank/Optimize","Generate plus\nRank/Optimize");d.rectangle((x,y,x+cell_w,y+90),fill="#173b4f")
        for j,p in enumerate(label.split("\n")):
            box=d.textbbox((0,0),p,font=f(18,True));d.text((x+(cell_w-(box[2]-box[0]))/2,y+25+j*23),p,font=f(18,True),fill="white");svg.append(f'<text x="{x+cell_w/2}" y="{y+44+j*23}" text-anchor="middle" font-family="Arial" font-size="18" font-weight="700" fill="#fff">{escape(p)}</text>')
    y+=90;current=None
    band_labels={"Band I":"Band I | AI capability with verified early-design AI context (8)","Band II":"Band II | AI capability without verified early-design AI context (7)","Band III":"Band III | no verified A/B Target/Output capability (4)"}
    for attr in attrs:
        band=bands[attr]
        if band!=current:
            current=band;d.rectangle((left,y,left+attr_w+6*cell_w,y+band_h),fill="#dcebf0");d.text((left+18,y+17),band_labels[band],font=f(21,True),fill="#173b4f");svg.append(f'<rect x="{left}" y="{y}" width="{attr_w+6*cell_w}" height="{band_h}" fill="#dcebf0"/><text x="{left+18}" y="{y+42}" font-family="Arial" font-size="21" font-weight="700" fill="#173b4f">{escape(band_labels[band])}</text>');y+=band_h
        d.rectangle((left,y,left+attr_w,y+row_h),fill="white",outline="#d9d9d9");d.text((left+18,y+22),attr,font=f(20,True),fill="#1f2933");svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="{row_h}" fill="#fff" stroke="#d9d9d9"/><text x="{left+18}" y="{y+50}" font-family="Arial" font-size="20" font-weight="700" fill="#1f2933">{escape(str(attr))}</text>')
        for i,fn in enumerate(FUNCTIONS):
            r=by[attr][fn];grade=r["grade"];x=left+attr_w+i*cell_w;shown=grade if grade in {"A","B"} else "no A/B";fill=COLORS[grade] if grade in {"A","B"} else "#f1f3f4";text="white" if grade=="A" else "#1f2933"
            d.rectangle((x,y,x+cell_w,y+row_h),fill=fill,outline="white");box=d.textbbox((0,0),shown,font=f(20,grade in {"A","B"}));d.text((x+(cell_w-(box[2]-box[0]))/2,y+22),shown,font=f(20,grade in {"A","B"}),fill=text);svg.append(f'<rect x="{x}" y="{y}" width="{cell_w}" height="{row_h}" fill="{fill}" stroke="#fff"/><text x="{x+cell_w/2}" y="{y+50}" text-anchor="middle" font-family="Arial" font-size="20" font-weight="700" fill="{text}">{escape(shown)}</text>')
            if grade in {"A","B"} and str(r["early_design_marker"]).strip().lower() in {"yes","y","true","1"}:
                d.ellipse((x+cell_w-52,y+14,x+cell_w-14,y+52),fill="white",outline="#173b4f");d.text((x+cell_w-46,y+23),"ED",font=f(13,True),fill="#173b4f");svg.append(f'<circle cx="{x+cell_w-33}" cy="{y+33}" r="19" fill="#fff" stroke="#173b4f"/><text x="{x+cell_w-33}" y="{y+38}" text-anchor="middle" font-family="Arial" font-size="13" font-weight="700" fill="#173b4f">ED</text>')
        y+=row_h
    note="Source-grounded readiness was established independently. ED marks an explicit or project-status-defensible early-design AI context for that A/B cell; it does not establish prospective availability of every required input."
    d.text((left,y+55),note,font=f(17),fill="#5d6871");svg.append(f'<text x="{left}" y="{y+75}" font-family="Arial" font-size="17" fill="#5d6871">{escape(note)}</text></svg>');SVG.write_text("".join(svg),encoding="utf-8");image.save(PNG,dpi=(300,300))

if __name__=="__main__":
    from collections import Counter
    main()
