from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data_excel" / "07_Figure_and_Table_Source_Data.xlsx"
SVG = ROOT / "figures" / "Figure_S1_Information_Source_Distribution.svg"
PNG = ROOT / "figures" / "Figure_S1_Information_Source_Distribution.png"
CATEGORIES = ["Drawing/CAD", "BIM/IFC", "Specification/BoQ", "Product Metadata", "EPD/LCA Database", "Existing-building Inventory", "Scenario/Judgment", "Not Reported"]


def font(size, bold=False):
    for name in (["arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "DejaVuSans.ttf"]):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            pass
    return ImageFont.load_default()


def mix(a, b, t):
    ar = tuple(int(a[i:i + 2], 16) for i in (1, 3, 5)); br = tuple(int(b[i:i + 2], 16) for i in (1, 3, 5))
    return "#" + "".join(f"{round(x + (y-x)*t):02x}" for x, y in zip(ar, br))


def main():
    ws = load_workbook(SOURCE, read_only=True, data_only=True)["Figure_S1"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(v) for v in rows[2]]
    data = [dict(zip(headers, row)) for row in rows[3:] if row[1]]
    assert len(data) == 48 and len({row["Normalized_Attribute"] for row in data}) == 48
    data.sort(key=lambda row: (str(row["Navigation_Family"]), str(row["Normalized_Attribute"])))
    max_count = max(int(row[c] or 0) for row in data for c in CATEGORIES)
    width, height = 3400, 4680
    left, top, attr_w, cell_w, row_h, family_h = 125, 360, 930, 285, 62, 56
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">', '<rect width="100%" height="100%" fill="#ffffff"/>']
    image = Image.new("RGB", (width, height), "white"); draw = ImageDraw.Draw(image)
    svg.append('<text x="125" y="105" font-family="Arial" font-size="52" font-weight="700" fill="#173b4f">Information-source distribution across all 48 normalized circularity attributes</text>')
    draw.text((125, 55), "Information-source distribution across all 48 normalized circularity attributes", font=font(52, True), fill="#173b4f")
    y = top
    header_fill = "#173b4f"
    draw.rectangle((left, y, left + attr_w, y + 90), fill=header_fill)
    svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="90" fill="{header_fill}"/>')
    draw.text((left + 18, y + 27), "Normalized attribute", font=font(24, True), fill="white")
    svg.append(f'<text x="{left+18}" y="{y+55}" font-family="Arial" font-size="24" font-weight="700" fill="#ffffff">Normalized attribute</text>')
    for i, category in enumerate(CATEGORIES):
        x = left + attr_w + i * cell_w
        draw.rectangle((x, y, x + cell_w, y + 90), fill=header_fill)
        label = category.replace("/", "/\n").replace(" ", "\n", 1) if len(category) > 14 else category
        parts = label.split("\n")
        for j, part in enumerate(parts):
            box = draw.textbbox((0, 0), part, font=font(19, True)); draw.text((x + (cell_w - (box[2]-box[0]))/2, y + 18 + j*23), part, font=font(19, True), fill="white")
            svg.append(f'<text x="{x+cell_w/2}" y="{y+35+j*23}" text-anchor="middle" font-family="Arial" font-size="19" font-weight="700" fill="#ffffff">{escape(part)}</text>')
    y += 90
    current_family = None
    for row in data:
        family = row["Navigation_Family"]
        if family != current_family:
            current_family = family
            draw.rectangle((left, y, left + attr_w + len(CATEGORIES)*cell_w, y + family_h), fill="#dcebf0")
            draw.text((left + 18, y + 13), family, font=font(21, True), fill="#173b4f")
            svg.append(f'<rect x="{left}" y="{y}" width="{attr_w+len(CATEGORIES)*cell_w}" height="{family_h}" fill="#dcebf0"/>')
            svg.append(f'<text x="{left+18}" y="{y+37}" font-family="Arial" font-size="21" font-weight="700" fill="#173b4f">{escape(str(family))}</text>')
            y += family_h
        draw.rectangle((left, y, left + attr_w, y + row_h), fill="white", outline="#d9d9d9")
        draw.text((left + 18, y + 16), str(row["Normalized_Attribute"]), font=font(19), fill="#1f2933")
        svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="{row_h}" fill="#ffffff" stroke="#d9d9d9"/>')
        svg.append(f'<text x="{left+18}" y="{y+39}" font-family="Arial" font-size="19" fill="#1f2933">{escape(str(row["Normalized_Attribute"]))}</text>')
        for i, category in enumerate(CATEGORIES):
            value = int(row[category] or 0); intensity = 0 if value == 0 else __import__("math").log1p(value)/__import__("math").log1p(max_count)
            low, high = ("#f6f7f8", "#7b8790") if category == "Not Reported" else ("#f3f8f8", "#24747a")
            fill = mix(low, high, intensity); text_color = "white" if intensity >= .66 else "#1f2933"
            x = left + attr_w + i * cell_w
            draw.rectangle((x, y, x + cell_w, y + row_h), fill=fill, outline="white")
            value_text = str(value); box = draw.textbbox((0, 0), value_text, font=font(19, True)); draw.text((x+(cell_w-(box[2]-box[0]))/2, y+16), value_text, font=font(19, True), fill=text_color)
            svg.append(f'<rect x="{x}" y="{y}" width="{cell_w}" height="{row_h}" fill="{fill}" stroke="#ffffff"/>')
            svg.append(f'<text x="{x+cell_w/2}" y="{y+39}" text-anchor="middle" font-family="Arial" font-size="19" font-weight="700" fill="{text_color}">{value}</text>')
        y += row_h
    note = "Counts are unique primary-review records with explicit broad source-group coding. Source groups are non-exclusive; missing source coding remains Not Reported."
    draw.text((left, y + 45), note, font=font(18), fill="#5d6871")
    svg.append(f'<text x="{left}" y="{y+65}" font-family="Arial" font-size="18" fill="#5d6871">{escape(note)}</text>')
    svg.append("</svg>")
    SVG.write_text("".join(svg), encoding="utf-8"); image.save(PNG, dpi=(300, 300))


if __name__ == "__main__":
    main()
