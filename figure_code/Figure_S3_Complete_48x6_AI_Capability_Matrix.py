from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data_excel" / "07_Figure_and_Table_Source_Data.xlsx"
SVG = ROOT / "figures" / "Figure_S3_Complete_48x6_AI_Capability_Matrix.svg"
PNG = ROOT / "figures" / "Figure_S3_Complete_48x6_AI_Capability_Matrix.png"
FUNCTIONS = ["Extract/Retrieve", "Infer/Complete", "Predict/Assess", "Diagnose", "Recommend", "Generate plus Rank/Optimize"]
COLORS = {"A":"#24747a", "B":"#d38b1f", "C":"#d8e8ee", "D":"#f1f3f4"}

def f(size,bold=False):
    for name in (["arialbd.ttf","DejaVuSans-Bold.ttf"] if bold else ["arial.ttf","DejaVuSans.ttf"]):
        try:return ImageFont.truetype(name,size)
        except OSError:pass
    return ImageFont.load_default()

def main():
    ws=load_workbook(SOURCE,read_only=True,data_only=True)["Figure_S3"]; rows=list(ws.iter_rows(values_only=True)); h=[str(v) for v in rows[2]]; data=[dict(zip(h,r)) for r in rows[3:] if r[0]]
    assert len(data)==288
    grades={g:sum(1 for r in data if r["grade"]==g) for g in "ABCD"}; assert grades=={"A":63,"B":9,"C":34,"D":182}
    attrs=[]; by={}
    for r in data:
        if r["attribute"] not in by: attrs.append((r["family"],r["attribute"],r["core_status"])); by[r["attribute"]]={}
        by[r["attribute"]][r["ai_function"]]=r
    assert len(attrs)==48
    attrs.sort(key=lambda item:(str(item[0]),str(item[1])))
    width,height=3400,4520; image=Image.new("RGB",(width,height),"white"); d=ImageDraw.Draw(image); svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#fff"/>']
    d.text((125,45),"Complete 48 × 6 AI capability matrix",font=f(56,True),fill="#173b4f"); svg.append('<text x="125" y="105" font-family="Arial" font-size="56" font-weight="700" fill="#173b4f">Complete 48 × 6 AI capability matrix</text>')
    left,top,attr_w,cell_w,row_h,fam_h=125,190,1120,335,58,52; y=top
    d.rectangle((left,y,left+attr_w,y+85),fill="#173b4f"); d.text((left+18,y+28),"Normalized attribute",font=f(22,True),fill="white"); svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="85" fill="#173b4f"/><text x="{left+18}" y="{y+54}" font-family="Arial" font-size="22" font-weight="700" fill="#fff">Normalized attribute</text>')
    for i,fn in enumerate(FUNCTIONS):
        x=left+attr_w+i*cell_w; label=fn.replace("Generate plus Rank/Optimize","Generate plus\nRank/Optimize")
        d.rectangle((x,y,x+cell_w,y+85),fill="#173b4f"); parts=label.split("\n")
        for j,p in enumerate(parts):
            box=d.textbbox((0,0),p,font=f(18,True)); d.text((x+(cell_w-(box[2]-box[0]))/2,y+23+j*23),p,font=f(18,True),fill="white")
            svg.append(f'<text x="{x+cell_w/2}" y="{y+42+j*23}" text-anchor="middle" font-family="Arial" font-size="18" font-weight="700" fill="#fff">{escape(p)}</text>')
    y+=85; current=None
    for family,attr,core in attrs:
        if family!=current:
            current=family; d.rectangle((left,y,left+attr_w+6*cell_w,y+fam_h),fill="#dcebf0"); d.text((left+15,y+13),family,font=f(19,True),fill="#173b4f"); svg.append(f'<rect x="{left}" y="{y}" width="{attr_w+6*cell_w}" height="{fam_h}" fill="#dcebf0"/><text x="{left+15}" y="{y+34}" font-family="Arial" font-size="19" font-weight="700" fill="#173b4f">{escape(str(family))}</text>'); y+=fam_h
        d.rectangle((left,y,left+attr_w,y+row_h),fill="white",outline="#d9d9d9"); d.text((left+18,y+15),attr,font=f(18),fill="#1f2933");
        if "Core" in str(core): d.ellipse((left+8,y+23,left+18,y+33),fill="#173b4f")
        svg.append(f'<rect x="{left}" y="{y}" width="{attr_w}" height="{row_h}" fill="#fff" stroke="#d9d9d9"/><text x="{left+18}" y="{y+37}" font-family="Arial" font-size="18" fill="#1f2933">{escape(str(attr))}</text>')
        for i,fn in enumerate(FUNCTIONS):
            r=by[attr][fn]; grade=r["grade"]; x=left+attr_w+i*cell_w; fill=COLORS[grade]; text="white" if grade=="A" else "#1f2933"
            d.rectangle((x,y,x+cell_w,y+row_h),fill=fill,outline="white"); box=d.textbbox((0,0),grade,font=f(19,True)); d.text((x+(cell_w-(box[2]-box[0]))/2,y+15),grade,font=f(19,True),fill=text)
            svg.append(f'<rect x="{x}" y="{y}" width="{cell_w}" height="{row_h}" fill="{fill}" stroke="#fff"/><text x="{x+cell_w/2}" y="{y+37}" text-anchor="middle" font-family="Arial" font-size="19" font-weight="700" fill="{text}">{grade}</text>')
            if grade in {"A","B"} and str(r["early_design_marker"]).strip().lower() in {"yes","y","true","1"}:
                d.ellipse((x+cell_w-29,y+7,x+cell_w-13,y+23),fill="white",outline="#173b4f"); svg.append(f'<circle cx="{x+cell_w-21}" cy="{y+15}" r="8" fill="#fff" stroke="#173b4f"/>')
        y+=row_h
    note="A/B are demonstrated Target/Output evidence; C is a supporting-role plausibility basis; D is no qualifying evidence in the selected corpus. D is not technical impossibility."
    d.text((left,y+55),note,font=f(17),fill="#5d6871"); svg.append(f'<text x="{left}" y="{y+75}" font-family="Arial" font-size="17" fill="#5d6871">{escape(note)}</text></svg>')
    SVG.write_text("".join(svg),encoding="utf-8"); image.save(PNG,dpi=(300,300))

if __name__=="__main__":main()
