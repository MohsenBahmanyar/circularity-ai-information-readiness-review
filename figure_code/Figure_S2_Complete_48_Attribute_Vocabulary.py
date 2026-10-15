from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data_excel" / "07_Figure_and_Table_Source_Data.xlsx"
SVG = ROOT / "figures" / "Figure_S2_Complete_48_Attribute_Vocabulary.svg"
PNG = ROOT / "figures" / "Figure_S2_Complete_48_Attribute_Vocabulary.png"

def f(size, bold=False):
    for name in (["arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold else ["arial.ttf", "DejaVuSans.ttf"]):
        try: return ImageFont.truetype(name, size)
        except OSError: pass
    return ImageFont.load_default()

def short(text, n):
    text = str(text or "")
    return text if len(text) <= n else text[:n-1] + "…"

def main():
    ws = load_workbook(SOURCE, read_only=True, data_only=True)["Figure_S2"]
    rows = list(ws.iter_rows(values_only=True)); headers = [str(v) for v in rows[2]]
    data = [dict(zip(headers, r)) for r in rows[3:] if r[0]]
    assert len(data) == 48 and len({r["normalized_attribute"] for r in data}) == 48
    data.sort(key=lambda row: (str(row["navigation_family"]), str(row["normalized_attribute"])))
    width, height = 4200, 3140; image = Image.new("RGB", (width, height), "white"); d = ImageDraw.Draw(image)
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><rect width="100%" height="100%" fill="#fff"/>']
    title = "Complete 48-attribute circularity-information vocabulary"
    d.text((120, 50), title, font=f(54, True), fill="#173b4f"); svg.append(f'<text x="120" y="105" font-family="Arial" font-size="54" font-weight="700" fill="#173b4f">{title}</text>')
    cols = [("navigation_family", "Navigation family", 760), ("normalized_attribute", "Normalized attribute", 670), ("requirement_frequency_n", "n", 150), ("core_status", "Readiness", 430), ("information_meaning", "Information meaning", 1130), ("representative_original_terminology", "Representative source terminology", 850)]
    x0, y, row_h = 100, 190, 55; header_h = 78; x = x0
    for key, label, w in cols:
        d.rectangle((x, y, x+w, y+header_h), fill="#173b4f"); d.text((x+14, y+24), label, font=f(20, True), fill="white")
        svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{header_h}" fill="#173b4f"/><text x="{x+14}" y="{y+48}" font-family="Arial" font-size="20" font-weight="700" fill="#fff">{escape(label)}</text>'); x += w
    y += header_h
    for i, row in enumerate(data):
        fill = "#ffffff" if i % 2 == 0 else "#f3f6f7"; x = x0
        for key, label, w in cols:
            value = row.get(key, "")
            if key == "core_status": value = "Source-grounded Core" if str(value) == "Source-grounded Core" else "Unresolved readiness"
            max_chars = {"navigation_family": 34, "normalized_attribute": 28, "information_meaning": 68, "representative_original_terminology": 52}.get(key, 20)
            value = short(value, max_chars)
            d.rectangle((x, y, x+w, y+row_h), fill=fill, outline="#d9d9d9"); color = "#1f2933" if key != "core_status" else ("#17645f" if "Core" in value else "#a56b16")
            d.text((x+12, y+16), value, font=f(17, key == "core_status"), fill=color)
            svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{row_h}" fill="{fill}" stroke="#d9d9d9"/><text x="{x+12}" y="{y+36}" font-family="Arial" font-size="17" fill="{color}">{escape(str(value))}</text>'); x += w
        y += row_h
    svg.append("</svg>"); SVG.write_text("".join(svg), encoding="utf-8"); image.save(PNG, dpi=(300, 300))

if __name__ == "__main__": main()
